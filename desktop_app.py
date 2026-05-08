#!/usr/bin/env python3
import atexit
import sys
import json
import os
import re
import platform
import subprocess
import time
import traceback
import faulthandler
import html
from datetime import datetime
BACKEND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

import requests
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QComboBox, QProgressBar,
    QTextEdit, QMessageBox, QFrame, QFileDialog, QScrollArea,
    QStackedWidget, QSizePolicy, QListWidget, QListWidgetItem, QAbstractItemView
)
from PyQt6.QtWidgets import QDialog, QPlainTextEdit
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QSize, QSettings
from PyQt6.QtGui import (
    QFont,
    QColor,
    QPalette,
    QIcon,
    QTextCursor,
    QTextBlockFormat,
    QCloseEvent,
    QKeySequence,
    QShortcut,
)

from app.services.prompt_validation import validate_rewrite_prompt

BACKEND_HEALTH_PATH = "/health"
BUILTIN_DEFAULT_REWRITE_PROMPT = (
    "保留原意和事实，不删关键信息，改写为更有节奏和可读性的中文内容。"
)
DEFAULT_BACKEND_CANDIDATES = [
    "http://127.0.0.1:8000",
    "http://localhost:8000",
    "http://127.0.0.1:8002",
    "http://localhost:8002",
]
_RESOLVED_BACKEND_URL = None
CRASH_LOG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "desktop_crash.log")
_CRASH_FILE_HANDLE = None
BUILTIN_STYLE_KEY = "__builtin_default__"


def _append_crash_log(title: str, detail: str) -> None:
    try:
        with open(CRASH_LOG_PATH, "a", encoding="utf-8") as handle:
            handle.write(f"\n[{datetime.now().isoformat(timespec='seconds')}] {title}\n")
            handle.write(detail.rstrip() + "\n")
    except Exception:
        pass


def _install_global_exception_hook() -> None:
    original_hook = sys.excepthook

    def _hook(exc_type, exc_value, exc_traceback):
        formatted = "".join(traceback.format_exception(exc_type, exc_value, exc_traceback))
        _append_crash_log("Unhandled exception", formatted)
        try:
            QMessageBox.critical(
                None,
                "程序异常",
                "GUI 遇到未处理异常，详情已写入 desktop_crash.log。",
            )
        except Exception:
            pass
        original_hook(exc_type, exc_value, exc_traceback)

    sys.excepthook = _hook


def _install_fault_handler() -> None:
    global _CRASH_FILE_HANDLE
    try:
        if _CRASH_FILE_HANDLE is None:
            _CRASH_FILE_HANDLE = open(CRASH_LOG_PATH, "a", encoding="utf-8")
            atexit.register(_CRASH_FILE_HANDLE.close)
        faulthandler.enable(file=_CRASH_FILE_HANDLE, all_threads=True)
    except Exception:
        pass


def _candidate_backend_urls():
    configured = (
        os.getenv("DESKTOP_BACKEND_URL")
        or os.getenv("NEXT_PUBLIC_API_BASE_URL")
        or ""
    ).strip().rstrip("/")
    if configured:
        return [configured, *[url for url in DEFAULT_BACKEND_CANDIDATES if url != configured]]
    return list(DEFAULT_BACKEND_CANDIDATES)


def resolve_backend_base_url(force_probe=False):
    global _RESOLVED_BACKEND_URL

    if _RESOLVED_BACKEND_URL and not force_probe:
        return _RESOLVED_BACKEND_URL

    for base_url in _candidate_backend_urls():
        try:
            response = requests.get(f"{base_url}{BACKEND_HEALTH_PATH}", timeout=2)
            if response.status_code != 200:
                continue
            try:
                payload = response.json()
            except Exception:
                continue
            if isinstance(payload, dict) and str(payload.get("status", "")).strip().lower() == "ok":
                _RESOLVED_BACKEND_URL = base_url
                return base_url
        except Exception:
            continue

    raise RuntimeError(
        "Cannot connect to backend API. Start backend on port 8000 or 8002, "
        "or set DESKTOP_BACKEND_URL."
    )


def bootstrap_local_backend_daemon() -> tuple[bool, str, str]:
    root_dir = os.path.dirname(os.path.abspath(__file__))
    script_path = os.path.join(root_dir, "start-backend.sh")
    if not os.path.isfile(script_path):
        return False, "", f"未找到后端启动脚本: {script_path}"

    try:
        result = subprocess.run(
            ["/bin/bash", script_path, "--daemon"],
            cwd=root_dir,
            capture_output=True,
            text=True,
            timeout=45,
        )
    except Exception as exc:
        return False, "", f"自动启动后端失败: {exc}"

    output = "\n".join(
        part.strip()
        for part in [result.stdout, result.stderr]
        if part and part.strip()
    ).strip()
    if result.returncode != 0:
        return False, "", output or "start-backend.sh --daemon 执行失败"

    for _ in range(20):
        try:
            backend_url = resolve_backend_base_url(force_probe=True)
            return True, backend_url, output
        except Exception:
            time.sleep(0.5)

    return False, "", output or "后端启动后仍未通过健康检查"


class WorkerThread(QThread):
    progress = pyqtSignal(int, str)
    result = pyqtSignal(dict)
    error = pyqtSignal(str)

    def __init__(
        self,
        backend_url,
        url,
        provider,
        source_mode,
        api_key,
        model,
        skill_prompt="",
    ):
        super().__init__()
        self.backend_url = backend_url
        self.url = url
        self.provider = provider
        self.source_mode = source_mode
        self.api_key = api_key
        self.model = model
        self.skill_prompt = skill_prompt

    def run(self):
        try:
            if self.source_mode == "subtitle_first":
                self.progress.emit(10, "正在处理：优先字幕路径...")
            else:
                self.progress.emit(10, "正在处理：音频 + Whisper 路径...")

            translation_config = {"provider": self.provider}
            if self.api_key:
                translation_config["api_key"] = self.api_key
            if self.model:
                translation_config["model"] = self.model

            resp = requests.post(
                f"{self.backend_url}/api/jobs/run",
                json={
                    "url": self.url,
                    "source_mode": self.source_mode,
                    "translation_config": translation_config
                },
                timeout=600
            )

            if resp.status_code >= 400:
                try:
                    detail = resp.json().get("detail", resp.text)
                except Exception:
                    detail = resp.text
                self.error.emit(str(detail or "jobs/run failed"))
                return

            data = resp.json()
            if not data.get("ok"):
                self.error.emit(data.get("detail", "Unknown error"))
                return

            source_type = str(data.get("source_type", "")).strip()
            if source_type == "captions":
                self.progress.emit(45, "已获取字幕，正在翻译并改写...")
            elif source_type == "audio":
                self.progress.emit(45, "已获取音频，正在转录、翻译并改写...")

            translation = data.get("translation_zh", {})
            zh_segments = translation.get("segments", [])
            zh_text = "\n".join(
                str(item.get("translated_text", "")).strip()
                for item in zh_segments
                if str(item.get("translated_text", "")).strip()
            ).strip()

            if zh_text:
                self.progress.emit(75, "正在改写中文内容...")
                rewrite_payload = {
                    "source_text": zh_text,
                    "translation_config": translation_config,
                    "rewrite_focus": self.skill_prompt,
                }
                content_context_id = str(data.get("content_context_id", "")).strip()
                if content_context_id:
                    rewrite_payload["content_context_id"] = content_context_id
                rewrite_resp = requests.post(
                    f"{self.backend_url}/api/content-rewrite",
                    json=rewrite_payload,
                    timeout=600,
                )
                if rewrite_resp.status_code >= 400:
                    try:
                        detail = rewrite_resp.json().get("detail", rewrite_resp.text)
                    except Exception:
                        detail = rewrite_resp.text
                    self.error.emit(str(detail or "content-rewrite failed"))
                    return
                rewrite_data = rewrite_resp.json()
                data["rewrite_zh"] = str(rewrite_data.get("rewritten_text", "")).strip()
                data["rewrite_focus"] = self.skill_prompt
                data["rewrite_source_text"] = zh_text
                data["rewrite_quality_issues"] = list(
                    rewrite_data.get("quality_issues", []) or []
                )
            else:
                data["rewrite_zh"] = ""
                data["rewrite_focus"] = self.skill_prompt
                data["rewrite_source_text"] = zh_text
                data["rewrite_quality_issues"] = []

            self.progress.emit(100, "完成!")
            self.result.emit(data)

        except requests.exceptions.Timeout:
            self.error.emit("请求超时，视频可能太长或后端忙碌")
        except requests.exceptions.ConnectionError as e:
            msg = str(e or "")
            lowered = msg.lower()
            if "remotedisconnected" in lowered or "connection aborted" in lowered:
                self.error.emit("上游模型连接被中断（RemoteDisconnected），请重试；若反复出现，请检查 API Key / base_url / 模型可用性。")
            else:
                self.error.emit(msg or "连接失败，请检查网络或后端状态")
        except Exception as e:
            self.error.emit(str(e))


class RewriteWorker(QThread):
    progress = pyqtSignal(int, str)
    result = pyqtSignal(dict)
    error = pyqtSignal(str)

    def __init__(
        self,
        backend_url,
        source_text,
        rewrite_focus,
        translation_config,
        content_context_id="",
    ):
        super().__init__()
        self.backend_url = backend_url
        self.source_text = source_text
        self.rewrite_focus = rewrite_focus
        self.translation_config = translation_config
        self.content_context_id = content_context_id

    def run(self):
        try:
            self.progress.emit(20, "正在重新改写内容...")
            payload = {
                "source_text": self.source_text,
                "translation_config": self.translation_config,
                "rewrite_focus": self.rewrite_focus,
            }
            if self.content_context_id:
                payload["content_context_id"] = self.content_context_id
            resp = requests.post(
                f"{self.backend_url}/api/content-rewrite",
                json=payload,
                timeout=600,
            )
            if resp.status_code >= 400:
                try:
                    detail = resp.json().get("detail", resp.text)
                except Exception:
                    detail = resp.text
                self.error.emit(str(detail or "content-rewrite failed"))
                return

            rewrite_data = resp.json()
            self.progress.emit(100, "重写完成")
            self.result.emit(
                {
                    "rewrite_zh": str(rewrite_data.get("rewritten_text", "")).strip(),
                    "rewrite_quality_issues": list(
                        rewrite_data.get("quality_issues", []) or []
                    ),
                    "rewrite_focus": self.rewrite_focus,
                    "rewrite_source_text": self.source_text,
                    "rewrite_provider": str(rewrite_data.get("provider", "")).strip(),
                    "rewrite_model": str(rewrite_data.get("model", "")).strip(),
                    "content_context_id": self.content_context_id,
                }
            )
        except requests.exceptions.Timeout:
            self.error.emit("重新改写请求超时，请稍后重试")
        except requests.exceptions.ConnectionError as e:
            msg = str(e or "")
            lowered = msg.lower()
            if "remotedisconnected" in lowered or "connection aborted" in lowered:
                self.error.emit("上游模型连接被中断（RemoteDisconnected），请重试；若反复出现，请检查 API Key / base_url / 模型可用性。")
            else:
                self.error.emit(msg or "重新改写连接失败")
        except Exception as e:
            self.error.emit(str(e))


class ChatWorker(QThread):
    response = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(
        self,
        backend_url,
        question,
        content_context_id,
        transcript_en,
        translation_zh,
        chat_config,
    ):
        super().__init__()
        self.backend_url = backend_url
        self.question = question
        self.content_context_id = content_context_id
        self.transcript_en = transcript_en
        self.translation_zh = translation_zh
        self.chat_config = chat_config

    def run(self):
        try:
            payload = {"question": self.question}
            if self.content_context_id:
                payload["content_context_id"] = self.content_context_id
            else:
                payload["transcript_en"] = self.transcript_en
                payload["translation_zh"] = self.translation_zh
            if self.chat_config:
                payload["chat_config"] = self.chat_config

            resp = requests.post(
                f"{self.backend_url}/api/content-chat",
                json=payload,
                timeout=60
            )
            if resp.status_code >= 400:
                try:
                    detail = resp.json().get("detail", resp.text)
                except Exception:
                    detail = resp.text
                self.error.emit(str(detail or "content-chat failed"))
                return
            data = resp.json()
            if data.get("ok"):
                self.response.emit(data.get("answer", "无回答"))
            else:
                self.error.emit(data.get("detail", "Chat failed"))
        except Exception as e:
            self.error.emit(str(e))


class HistoryListWorker(QThread):
    result = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, backend_url, limit=20):
        super().__init__()
        self.backend_url = backend_url
        self.limit = limit

    def run(self):
        try:
            resp = requests.get(
                f"{self.backend_url}/api/session-history",
                params={"limit": self.limit},
                timeout=20,
            )
            if resp.status_code >= 400:
                try:
                    detail = resp.json().get("detail", resp.text)
                except Exception:
                    detail = resp.text
                self.error.emit(str(detail or "session-history list failed"))
                return
            data = resp.json()
            if data.get("ok"):
                self.result.emit(list(data.get("items", [])))
            else:
                self.error.emit(data.get("detail", "Failed to load history"))
        except Exception as e:
            self.error.emit(str(e))


class HistoryDetailWorker(QThread):
    result = pyqtSignal(dict)
    error = pyqtSignal(str)

    def __init__(self, backend_url, content_context_id):
        super().__init__()
        self.backend_url = backend_url
        self.content_context_id = content_context_id

    def run(self):
        try:
            resp = requests.get(
                f"{self.backend_url}/api/session-history/{self.content_context_id}",
                timeout=20,
            )
            if resp.status_code >= 400:
                try:
                    detail = resp.json().get("detail", resp.text)
                except Exception:
                    detail = resp.text
                self.error.emit(str(detail or "session-history detail failed"))
                return
            self.result.emit(resp.json())
        except Exception as e:
            self.error.emit(str(e))


class BackendBootstrapWorker(QThread):
    result = pyqtSignal(bool, str, str)

    def run(self):
        try:
            success, backend_url, info = bootstrap_local_backend_daemon()
            self.result.emit(success, backend_url, info)
        except Exception as exc:
            self.result.emit(False, "", str(exc))


class ChatTextEdit(QTextEdit):
    returnPressed = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptRichText(False)

    def keyPressEvent(self, event):
        if event.key() == Qt.Key.Key_Return and not event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
            self.returnPressed.emit()
            event.accept()
        else:
            super().keyPressEvent(event)


class YouTubeTranslatorApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("YouTube 翻译助手")
        self.setMinimumSize(1100, 750)
        self.resize(1280, 800)
        self.current_content = None
        self.backend_url = None
        self._backend_online = False
        self._backend_bootstrap_error = ""
        self._settings = None
        self.current_rewrite_raw_text = ""
        self.current_rewrite_source_text = ""
        self.current_rewrite_focus = ""
        self.current_rewrite_quality_issues = []
        self.last_exported_file_path = ""
        self.worker = None
        self.rewrite_worker = None
        self.chat_worker = None
        self.history_worker = None
        self.history_detail_worker = None
        self.backend_bootstrap_worker = None
        self.skills_dict = {}
        self.skill_prompt_validation = {}
        self.history_items = []
        self.history_search_text = ""
        self.current_history_id = ""
        self.chat_message_widgets = []
        self.chat_pending_message_widget = None
        self.chat_pending_message_label = None
        self.setup_ui()
        self.apply_light_theme()
        self._settings = QSettings("TranslationWritingWorkbench", "Desktop")
        self.load_skills()
        self._load_persisted_settings()
        self._install_keyboard_shortcuts()
        self.check_backend()

    def load_skills(self):
        skills_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "skills")
        self.skills_dict = {}
        if not os.path.exists(skills_dir):
            self.skill_prompt_validation = {}
            if hasattr(self, "skill_combo"):
                self._populate_skill_combo()
            return
        self.skill_prompt_validation = {}
        for file in os.listdir(skills_dir):
            if file.endswith(".md") or file.endswith(".txt"):
                path = os.path.join(skills_dir, file)
                name = os.path.splitext(file)[0]
                with open(path, "r", encoding="utf-8") as f:
                    prompt_text = f.read()
                    self.skills_dict[name] = prompt_text
                    self.skill_prompt_validation[name] = validate_rewrite_prompt(prompt_text)
        
        if hasattr(self, "skill_combo") and self.skills_dict:
            self._populate_skill_combo()

    def _populate_skill_combo(self):
        self.skill_combo.clear()
        builtin_validation = validate_rewrite_prompt(BUILTIN_DEFAULT_REWRITE_PROMPT)
        self.skill_combo.addItem("默认内置风格（晚点通用）", BUILTIN_STYLE_KEY)
        self.skill_prompt_validation[BUILTIN_STYLE_KEY] = builtin_validation
        self.skills_dict[BUILTIN_STYLE_KEY] = BUILTIN_DEFAULT_REWRITE_PROMPT
        for name, prompt_text in self.skills_dict.items():
            if name == BUILTIN_STYLE_KEY:
                continue
            validation = self.skill_prompt_validation.get(name) or validate_rewrite_prompt(prompt_text)
            if validation.errors:
                label = f"{name}（无效）"
            elif validation.has_transcript_placeholder:
                label = f"{name}（完整）"
            else:
                label = f"{name}（提示）"
            self.skill_combo.addItem(label, name)
        if self.skill_combo.count() == 1:
            self.skill_combo.setCurrentIndex(0)

    def _load_persisted_settings(self):
        if self._settings is None:
            return

        provider = str(self._settings.value("provider", "deepseek") or "deepseek").strip()
        model = str(self._settings.value("model", "deepseek-chat") or "deepseek-chat").strip()
        style_key = str(self._settings.value("style_key", BUILTIN_STYLE_KEY) or BUILTIN_STYLE_KEY).strip()
        api_key = str(self._settings.value("api_key", "") or "").strip()

        if provider:
            self.provider_combo.setCurrentText(provider)
        if model:
            self.model_input.setText(model)
        if api_key:
            self.api_key_input.setText(api_key)
        self._select_skill_by_key(style_key)

    def _persist_settings(self):
        if self._settings is None:
            return
        self._settings.setValue("provider", self.provider_combo.currentText())
        self._settings.setValue("model", self.model_input.text().strip())
        self._settings.setValue("style_key", self.skill_combo.currentData() or BUILTIN_STYLE_KEY)
        self._settings.setValue("api_key", self.api_key_input.text().strip())

    def _select_skill_by_key(self, skill_key: str):
        normalized_key = str(skill_key or "").strip() or BUILTIN_STYLE_KEY
        for index in range(self.skill_combo.count()):
            if str(self.skill_combo.itemData(index) or "").strip() == normalized_key:
                self.skill_combo.setCurrentIndex(index)
                return
        self.skill_combo.setCurrentIndex(0)

    def _install_keyboard_shortcuts(self):
        self.run_shortcut = QShortcut(QKeySequence("Ctrl+Return"), self)
        self.run_shortcut.activated.connect(self.start_processing)

        self.copy_shortcut = QShortcut(QKeySequence("Ctrl+Shift+C"), self)
        self.copy_shortcut.activated.connect(self.copy_rewrite_text)

        self.export_shortcut = QShortcut(QKeySequence("Ctrl+Shift+E"), self)
        self.export_shortcut.activated.connect(self.export_rewrite_text)

        self.chat_shortcut = QShortcut(QKeySequence("Ctrl+Shift+Return"), self)
        self.chat_shortcut.activated.connect(self.send_chat)

    def _current_rewrite_text(self) -> str:
        return str(self.current_rewrite_raw_text or "").strip()

    def _current_rewrite_source_text(self) -> str:
        text = str(self.current_rewrite_source_text or "").strip()
        if text:
            return text
        if not self.current_content:
            return ""
        translation = self.current_content.get("translation_zh", {})
        segments = translation.get("segments", []) if isinstance(translation, dict) else []
        return "\n".join(
            str(item.get("translated_text", "")).strip()
            for item in segments
            if str(item.get("translated_text", "")).strip()
        ).strip()

    def _current_rewrite_focus(self) -> str:
        focus = str(self.current_rewrite_focus or "").strip()
        if focus:
            return focus
        current_data = self.skill_combo.currentData()
        if not current_data:
            return ""
        key = str(current_data).strip()
        if key in self.skills_dict:
            return str(self.skills_dict.get(key, "")).strip()
        if key == BUILTIN_STYLE_KEY:
            return BUILTIN_DEFAULT_REWRITE_PROMPT
        return ""

    def _update_rewrite_quality_banner(self, issues: list[str] | None) -> None:
        issue_list = [str(item).strip() for item in (issues or []) if str(item).strip()]
        self.current_rewrite_quality_issues = issue_list
        if not issue_list:
            self.rewrite_issue_label.setVisible(False)
            self.rewrite_issue_label.setText("")
            return
        preview = "；".join(issue_list[:3])
        suffix = "…" if len(issue_list) > 3 else ""
        self.rewrite_issue_label.setText(f"改写提示：{preview}{suffix}")
        self.rewrite_issue_label.setVisible(True)

    def _set_status_text(self, text: str) -> None:
        message = str(text or "").strip()
        if hasattr(self, "progress_label"):
            self.progress_label.setText(message)
        if hasattr(self, "result_status_label"):
            self.result_status_label.setText(message)

    def _friendly_error_text(self, message: str) -> str:
        text = str(message or "").strip()
        lowered = text.lower()
        if "cookies are no longer valid" in lowered or "sign in to confirm you're not a bot" in lowered:
            return "YouTube Cookie 已失效或触发风控，请重新导出 cookies 后再试。"
        if "authentication fails" in lowered or "invalid api key" in lowered or "incorrect api key" in lowered:
            return "鉴权失败：当前 API Key 无效，请检查 provider、API Key 与 base_url 是否匹配。"
        if "remotedisconnected" in lowered or "connection aborted" in lowered:
            return "上游模型连接被中断（RemoteDisconnected），请重试；若反复出现，请检查 API Key / base_url / 模型可用性。"
        if "whisper" in lowered and ("error" in lowered or "failed" in lowered or "runtime" in lowered):
            return "Whisper 处理失败，请检查本地模型、设备配置或音频文件是否可用。"
        if "403" in text and "content-rewrite" in lowered:
            return "内容改写请求被上游拒绝，请检查模型服务、API Key 或 base_url。"
        return text

    def _current_video_title(self) -> str:
        if not self.current_content:
            return ""

        video = self.current_content.get("video")
        title = ""
        if isinstance(video, dict):
            title = str(video.get("title") or "").strip()
        if title:
            return title

        return str(self.current_content.get("video_title") or "").strip()

    def _current_session_id(self) -> str:
        if self.current_content:
            return str(self.current_content.get("content_context_id", "")).strip()
        return ""

    def _format_history_timestamp(self, value: str | None) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        return text.replace("T", " ")[:19]

    def _history_item_preview(self, item: dict) -> str:
        parts = []
        video_title = str(item.get("video_title") or "").strip()
        if video_title:
            parts.append(video_title)

        updated_at = self._format_history_timestamp(item.get("updated_at"))
        meta = []
        if updated_at:
            meta.append(updated_at)
        if item.get("source_type"):
            meta.append(str(item.get("source_type")))
        duration = item.get("video_duration_sec")
        if isinstance(duration, int) and duration > 0:
            meta.append(f"{duration}s")
        uploader = str(item.get("video_uploader") or "").strip()
        if uploader:
            meta.append(uploader)
        if item.get("chat_turn_count"):
            meta.append(f"{item.get('chat_turn_count')} 条问答")
        if item.get("has_rewrite"):
            meta.append("已改写")
        quality_issue_count = int(item.get("rewrite_quality_issue_count") or 0)
        if quality_issue_count:
            meta.append(f"{quality_issue_count} 个问题")
        if meta:
            parts.append(" · ".join(meta))

        rewritten = str(item.get("rewritten_preview") or "").strip()
        translation = str(item.get("translation_preview") or "").strip()
        preview = rewritten or translation or str(item.get("transcript_preview") or "").strip()
        if preview:
            parts.append(preview)

        return "\n".join(parts).strip() or "未命名会话"

    def _history_detail_to_current_content(self, detail: dict) -> dict:
        transcript_segments = [
            {
                "index": int(segment.get("index", 0)),
                "start": float(segment.get("start", 0.0)),
                "end": float(segment.get("end", 0.0)),
                "text": str(segment.get("text", "")).strip(),
                "speaker": segment.get("speaker"),
            }
            for segment in detail.get("transcript_en_segments", [])
            if str(segment.get("text", "")).strip()
        ]
        translation_segments = [
            {
                "index": int(segment.get("index", 0)),
                "start": float(segment.get("start", 0.0)),
                "end": float(segment.get("end", 0.0)),
                "source_text": str(segment.get("source_text", "")).strip(),
                "translated_text": str(segment.get("translated_text", "")).strip(),
            }
            for segment in detail.get("translation_zh_segments", [])
            if str(segment.get("translated_text", "")).strip()
        ]

        return {
            "ok": True,
            "video": {
                "video_id": detail.get("video_id", ""),
                "title": detail.get("video_title", "") or "",
                "thumbnail": detail.get("video_thumbnail", None),
                "duration_sec": detail.get("video_duration_sec", None),
                "uploader": detail.get("video_uploader", None),
            },
            "source_type": detail.get("source_type", "captions"),
            "content_context_id": detail.get("content_context_id", ""),
            "transcript_en": {
                "text": detail.get("transcript_en_text", ""),
                "segments": transcript_segments,
            },
            "translation_zh": {
                "segments": translation_segments,
            },
            "rewrite_zh": detail.get("rewritten_text", ""),
            "rewrite_source_text": detail.get("rewrite_source_text", ""),
            "rewrite_focus": detail.get("rewrite_focus", ""),
            "rewrite_quality_issues": detail.get("rewrite_quality_issues", []),
            "rewrite_provider": detail.get("rewrite_provider", ""),
            "rewrite_model": detail.get("rewrite_model", ""),
        }

    def refresh_history_sidebar(self):
        if not self.backend_url:
            return
        if self.history_worker and self.history_worker.isRunning():
            return

        self.history_status_label.setText("加载历史中...")
        self.history_refresh_button.setEnabled(False)
        self.history_list_widget.setEnabled(False)

        self.history_worker = HistoryListWorker(self.backend_url, limit=20)
        self.history_worker.result.connect(self.on_history_list_loaded)
        self.history_worker.error.connect(self.on_history_list_error)
        self.history_worker.finished.connect(self._on_history_list_finished)
        self.history_worker.finished.connect(self.history_worker.deleteLater)
        self.history_worker.start()

    def load_history_detail(self, content_context_id: str):
        if not self.backend_url:
            return
        normalized_id = str(content_context_id or "").strip()
        if not normalized_id:
            return
        if self.history_detail_worker and self.history_detail_worker.isRunning():
            return

        self.history_status_label.setText("正在打开历史会话...")
        self.history_refresh_button.setEnabled(False)
        self.history_list_widget.setEnabled(False)

        self.history_detail_worker = HistoryDetailWorker(self.backend_url, normalized_id)
        self.history_detail_worker.result.connect(self.on_history_detail_loaded)
        self.history_detail_worker.error.connect(self.on_history_detail_error)
        self.history_detail_worker.finished.connect(self._on_history_detail_finished)
        self.history_detail_worker.finished.connect(self.history_detail_worker.deleteLater)
        self.history_detail_worker.start()

    def on_history_list_loaded(self, items):
        self.history_items = list(items or [])
        self.render_history_list()

    def on_history_list_error(self, message):
        self.history_status_label.setText(f"历史加载失败：{message[:60]}")

    def _on_history_list_finished(self):
        self.history_refresh_button.setEnabled(True)
        self.history_list_widget.setEnabled(True)
        self.history_worker = None

    def on_history_search_changed(self, text):
        self.history_search_text = str(text or "").strip()
        self.render_history_list()

    def _history_matches_search(self, item: dict, query: str) -> bool:
        if not query:
            return True

        haystack_parts = [
            item.get("content_context_id"),
            item.get("video_id"),
            item.get("video_url"),
            item.get("video_title"),
            item.get("source_mode"),
            item.get("source_type"),
            item.get("transcript_preview"),
            item.get("translation_preview"),
            item.get("rewritten_preview"),
        ]
        haystack = " ".join(str(part or "") for part in haystack_parts).lower()
        return query.lower() in haystack

    def render_history_list(self):
        self.history_list_widget.clear()

        if not self.history_items:
            empty_item = QListWidgetItem("暂无历史记录\n处理过的视频会自动保存在这里。")
            empty_item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.history_list_widget.addItem(empty_item)
            self.history_status_label.setText("暂无历史记录。")
            return

        filtered_items = [
            item
            for item in self.history_items
            if self._history_matches_search(item, self.history_search_text)
        ]

        if not filtered_items:
            empty_item = QListWidgetItem("没有匹配到结果\n换个关键词再试试。")
            empty_item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.history_list_widget.addItem(empty_item)
            self.history_status_label.setText(
                f"已加载 {len(self.history_items)} 条历史会话，未匹配到结果。"
            )
            return

        for item in filtered_items:
            content_context_id = str(item.get("content_context_id", "")).strip()
            label = self._history_item_preview(item)
            list_item = QListWidgetItem(label)
            list_item.setData(Qt.ItemDataRole.UserRole, content_context_id)
            list_item.setSizeHint(QSize(250, 92))
            self.history_list_widget.addItem(list_item)

        self.history_status_label.setText(
            f"已加载 {len(self.history_items)} 条历史会话，匹配 {len(filtered_items)} 条。"
        )
        if self.current_history_id:
            self._select_history_item(self.current_history_id)

    def _select_history_item(self, content_context_id: str):
        normalized_id = str(content_context_id or "").strip()
        if not normalized_id:
            return

        for index in range(self.history_list_widget.count()):
            item = self.history_list_widget.item(index)
            if not item:
                continue
            item_id = str(item.data(Qt.ItemDataRole.UserRole) or "").strip()
            if item_id == normalized_id:
                self.history_list_widget.setCurrentItem(item)
                break

    def _on_history_item_clicked(self, item):
        content_context_id = str(
            item.data(Qt.ItemDataRole.UserRole) if item else ""
        ).strip()
        if not content_context_id:
            return
        self.load_history_detail(content_context_id)

    def on_history_detail_loaded(self, detail):
        content_context_id = str(detail.get("content_context_id", "")).strip()
        current_content = self._history_detail_to_current_content(detail)
        self.current_content = current_content
        self.last_exported_file_path = ""
        self.open_export_button.setEnabled(False)
        video_url = str(detail.get("video_url", "")).strip()
        if video_url:
            self.url_input.setText(video_url)
        rewritten_text = str(detail.get("rewritten_text", "")).strip()
        self.current_rewrite_raw_text = rewritten_text
        self.current_rewrite_source_text = str(detail.get("rewrite_source_text", "")).strip()
        self.current_rewrite_focus = str(detail.get("rewrite_focus", "") or "").strip()
        self._update_rewrite_quality_banner(detail.get("rewrite_quality_issues", []))
        translation_text = "\n\n".join(
            str(item.get("translated_text", "")).strip()
            for item in detail.get("translation_zh_segments", [])
            if str(item.get("translated_text", "")).strip()
        ).strip()
        if rewritten_text:
            self.trans_text.setMarkdown(rewritten_text)
        elif translation_text:
            self.trans_text.setMarkdown(translation_text)
        else:
            self.trans_text.setMarkdown("暂无可展示内容")

        has_rewrite = bool(self._current_rewrite_text())
        self.copy_button.setEnabled(has_rewrite)
        self.export_button.setEnabled(has_rewrite)
        self.rewrite_again_button.setEnabled(bool(rewritten_text or translation_text))
        self.chat_input.setEnabled(True)
        self.chat_button.setEnabled(True)
        self.current_history_id = content_context_id
        self._set_status_text("已加载历史会话")
        self.stack.setCurrentIndex(1)
        self.refresh_history_sidebar()

    def on_history_detail_error(self, message):
        friendly = self._friendly_error_text(message)
        self._set_status_text(f"历史打开失败: {friendly[:50]}")
        QMessageBox.warning(self, "历史会话", friendly)

    def _on_history_detail_finished(self):
        self.history_refresh_button.setEnabled(True)
        self.history_list_widget.setEnabled(True)
        self.history_detail_worker = None

    def _start_local_backend_daemon(self) -> tuple[bool, str]:
        started, backend_url, info = bootstrap_local_backend_daemon()
        if started:
            self.backend_url = backend_url
            return True, info
        return False, info

    def _set_backend_online_ui(self):
        self._backend_online = True

    def _set_backend_offline_ui(self):
        self._backend_online = False

    def apply_light_theme(self):
        self.setStyleSheet("""
            QMainWindow {
                background-color: #F8FAFD;
            }
            QWidget {
                background-color: #F8FAFD;
                color: #0F172A;
                font-family: "Google Sans", "PingFang SC", "Inter";
                font-size: 14px;
                letter-spacing: 0.2px;
            }
            #sidebar {
                background-color: #F8FAFD;
                border-right: none;
            }
            #sidebar_content {
                background-color: #F8FAFD;
            }
            QFrame {
                background-color: #FFFFFF;
                border-radius: 24px;
                border: none;
            }
            #rewrite_card {
                border-radius: 24px;
                background-color: #FFFFFF;
            }
            #rewrite_toolbar {
                background-color: transparent;
                border: none;
            }
            #rewrite_body {
                background-color: transparent;
                border: none;
            }
            #rewrite_action_btn {
                background-color: #F1F5F9;
                color: #334155;
                border: none;
                border-radius: 16px;
                padding: 8px 16px;
                font-size: 13px;
                font-weight: 600;
                min-height: 32px;
            }
            #rewrite_action_btn:hover {
                background-color: #E2E8F0;
            }
            #rewrite_action_btn:disabled {
                background-color: transparent;
                color: transparent;
            }
            #no_border_frame {
                background-color: transparent;
                border: none;
                border-radius: 0px;
            }
            QLabel {
                background-color: transparent;
                color: #0F172A;
            }
            QLineEdit {
                background-color: #FFFFFF;
                border: 1px solid #E2E8F0;
                border-radius: 12px;
                padding: 12px 16px;
                color: #0F172A;
                font-size: 14px;
                min-height: 24px;
            }
            QLineEdit:focus {
                border: 1px solid #5B6EFF;
                background-color: #F8FAFF;
            }
            QLineEdit::placeholder {
                color: #94A3B8;
            }
            QComboBox {
                background-color: #FFFFFF;
                border: 1px solid #E2E8F0;
                border-radius: 12px;
                padding: 10px 14px;
                color: #0F172A;
                font-size: 14px;
                min-height: 24px;
            }
            QComboBox:focus {
                border: 1px solid #5B6EFF;
            }
            QComboBox::drop-down {
                border: none;
                width: 24px;
            }
            QComboBox::down-arrow {
                image: none;
                border-left: 4px solid transparent;
                border-right: 4px solid transparent;
                border-top: 5px solid #86868B;
                width: 0px;
                height: 0px;
            }
            QPushButton {
                background-color: #5B6EFF;
                color: #FFFFFF;
                border: none;
                border-radius: 16px;
                padding: 12px 24px;
                font-weight: 600;
                font-size: 15px;
            }
            QPushButton:hover {
                background-color: #4B5DE6;
            }
            QPushButton:disabled {
                background-color: #E2E8F0;
                color: #94A3B8;
            }
            #chat_card {
                background-color: #FFFFFF;
                border: 1px solid #E2E8F0;
                border-radius: 24px;
            }
            #chat_composer {
                background-color: #FFFFFF;
                border: none;
                border-radius: 28px;
            }
            #chat_composer QTextEdit {
                background: transparent;
                border: none;
                padding: 4px 8px;
                font-size: 15px;
                color: #0F172A;
            }
            #chat_composer QTextEdit:focus {
                background-color: transparent;
            }
            #chat_send_btn {
                background-color: #101828;
                border-radius: 20px;
                min-width: 40px;
                max-width: 40px;
                min-height: 40px;
                max-height: 40px;
                padding: 0px;
                font-size: 18px;
                font-weight: 700;
            }
            #chat_send_btn:hover {
                background-color: #0B1220;
            }
            #chat_send_btn:disabled {
                background-color: #E2E8F0;
                color: #94A3B8;
            }
            #history_card {
                background-color: #FFFFFF;
                border: 1px solid #E2E8F0;
                border-radius: 24px;
            }
            #history_list {
                background-color: transparent;
                border: none;
                outline: none;
            }
            #history_list::item {
                background-color: #F8FAFC;
                border: 1px solid #E2E8F0;
                border-radius: 16px;
                padding: 12px 14px;
                color: #0F172A;
                margin: 0px;
                white-space: pre-wrap;
            }
            #history_list::item:selected {
                background-color: #EEF2FF;
                border-color: #A5B4FC;
            }
            QProgressBar {
                border: none;
                border-radius: 2px;
                background-color: #E2E8F0;
                text-align: center;
                color: transparent;
            }
            QProgressBar::chunk {
                background-color: #5B6EFF;
                border-radius: 2px;
            }
            QScrollArea {
                border: none;
                background-color: transparent;
            }
            QTextEdit {
                background-color: transparent;
                border: none;
                color: #1E293B;
                font-size: 16px;
                line-height: 1.8;
                padding: 0px;
            }
            QTextEdit::placeholder {
                color: #CBD5E1;
            }
            QScrollBar:vertical {
                background: transparent;
                width: 6px;
                border-radius: 3px;
            }
            QScrollBar::handle:vertical {
                background: #CBD5E1;
                border-radius: 3px;
                min-height: 24px;
            }
            QScrollBar::handle:vertical:hover {
                background: #94A3B8;
            }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
                height: 0px;
            }
        """)

    def setup_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        
        main_layout = QVBoxLayout(central)
        main_layout.setSpacing(0)
        main_layout.setContentsMargins(0, 0, 0, 0)
        
        self.stack = QStackedWidget(central)
        main_layout.addWidget(self.stack)
        
        self.input_page = self._create_input_page()
        self.result_page = self._create_result_page()
        
        self.stack.addWidget(self.input_page)
        self.stack.addWidget(self.result_page)
        self.stack.setCurrentIndex(0)

    def _create_input_page(self):
        page = QWidget()
        page.setStyleSheet("background-color: #F8FAFD;")
        layout = QVBoxLayout(page)
        layout.setContentsMargins(40, 40, 40, 40)
        
        # Center Title
        title_layout = QVBoxLayout()
        title_label = QLabel("YouTube 翻译助手")
        title_label.setStyleSheet("font-size: 32px; font-weight: 700; color: #0F172A; background: transparent;")
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title_layout.addWidget(title_label)
        
        layout.addStretch(2)
        layout.addLayout(title_layout)
        layout.addSpacing(40)
        
        # Input composer
        composer = QFrame()
        composer.setObjectName("chat_composer")
        composer.setMinimumHeight(60)
        composer.setMaximumWidth(800)
        composer.setStyleSheet("QFrame#chat_composer { border-radius: 30px; background-color: #FFFFFF; }")
        
        composer_layout = QHBoxLayout(composer)
        composer_layout.setContentsMargins(20, 10, 10, 10)
        composer_layout.setSpacing(12)
        
        self.url_input = QLineEdit()
        self.url_input.setPlaceholderText("粘贴 YouTube 链接...")
        self.url_input.setStyleSheet("border: none; background: transparent; font-size: 16px;")
        self.url_input.returnPressed.connect(self.start_processing)
        
        self.run_button = QPushButton("➤")
        self.run_button.setFixedSize(40, 40)
        self.run_button.setStyleSheet("""
            QPushButton {
                background-color: #101828;
                color: #FFFFFF;
                border-radius: 20px;
                font-size: 18px;
            }
            QPushButton:hover { background-color: #0B1220; }
            QPushButton:disabled { background-color: #E2E8F0; color: #94A3B8; }
        """)
        self.run_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.run_button.clicked.connect(self.start_processing)
        
        composer_layout.addWidget(self.url_input, 1)
        composer_layout.addWidget(self.run_button)
        
        center_layout = QHBoxLayout()
        center_layout.addStretch()
        center_layout.addWidget(composer, 1)
        center_layout.addStretch()
        
        layout.addLayout(center_layout)
        
        # Progress Area
        prog_layout = QVBoxLayout()
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setFixedHeight(4)
        self.progress_bar.setMaximumWidth(800)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setVisible(False)
        
        self.progress_label = QLabel("")
        self.progress_label.setStyleSheet("color: #86868B; font-size: 14px; background: transparent;")
        self.progress_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        
        prog_center = QHBoxLayout()
        prog_center.addStretch()
        prog_layout.addWidget(self.progress_bar, 0, Qt.AlignmentFlag.AlignHCenter)
        prog_layout.addWidget(self.progress_label, 0, Qt.AlignmentFlag.AlignHCenter)
        prog_center.addLayout(prog_layout)
        prog_center.addStretch()
        
        layout.addSpacing(20)
        layout.addLayout(prog_center)
        
        layout.addStretch(2)
        
        # Bottom left gear & Settings
        self.settings_panel = self._create_settings_card()
        self.settings_panel.setVisible(False)
        self.settings_panel.setMaximumWidth(400)
        
        self.gear_btn = QPushButton("⚙️")
        self.gear_btn.setFixedSize(44, 44)
        self.gear_btn.setStyleSheet("""
            QPushButton {
                background-color: #FFFFFF;
                color: #475569;
                border-radius: 22px;
                font-size: 24px;
                border: 1px solid #E2E8F0;
            }
            QPushButton:hover {
                background-color: #F1F5F9;
                border-color: #CBD5E1;
            }
        """)
        self.gear_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.gear_btn.clicked.connect(lambda: self.settings_panel.setVisible(not self.settings_panel.isVisible()))
        
        bottom_layout = QHBoxLayout()
        bottom_left = QVBoxLayout()
        bottom_left.addWidget(self.settings_panel)
        bottom_left.addWidget(self.gear_btn, 0, Qt.AlignmentFlag.AlignBottom | Qt.AlignmentFlag.AlignLeft)
        
        bottom_layout.addLayout(bottom_left)
        bottom_layout.addStretch(1)
        
        layout.addLayout(bottom_layout)
        
        return page

    def _create_result_page(self):
        page = QWidget()
        page.setStyleSheet("background-color: #F8FAFD;")
        layout = QHBoxLayout(page)
        layout.setContentsMargins(40, 40, 40, 40)
        layout.setSpacing(20)

        self.history_sidebar = self._create_history_sidebar()
        layout.addWidget(self.history_sidebar, 0)

        main_content = QWidget()
        main_layout = QVBoxLayout(main_content)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(20)

        # Header (Back button + Action buttons on the right)
        header_layout = QHBoxLayout()

        back_btn = QPushButton("← 返回并新开一个")
        back_btn.setStyleSheet("""
            QPushButton {
                background-color: #E2E8F0;
                color: #334155;
                border-radius: 16px;
                padding: 8px 16px;
                font-weight: bold;
                font-size: 14px;
            }
            QPushButton:hover { background-color: #CBD5E1; }
        """)
        back_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        back_btn.clicked.connect(self.go_to_input_page)
        header_layout.addWidget(back_btn)

        header_layout.addStretch()

        self.copy_button = QPushButton("复制内容")
        self.copy_button.setObjectName("rewrite_action_btn")
        self.copy_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.copy_button.setEnabled(False)
        self.copy_button.clicked.connect(self.copy_rewrite_text)
        header_layout.addWidget(self.copy_button)

        self.export_button = QPushButton("导出为...")
        self.export_button.setObjectName("rewrite_action_btn")
        self.export_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.export_button.setEnabled(False)
        self.export_button.clicked.connect(self.export_rewrite_text)
        header_layout.addWidget(self.export_button)

        self.rewrite_again_button = QPushButton("再改写")
        self.rewrite_again_button.setObjectName("rewrite_action_btn")
        self.rewrite_again_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.rewrite_again_button.setEnabled(False)
        self.rewrite_again_button.clicked.connect(self.rewrite_current_result)
        header_layout.addWidget(self.rewrite_again_button)

        self.open_export_button = QPushButton("打开位置")
        self.open_export_button.setObjectName("rewrite_action_btn")
        self.open_export_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.open_export_button.setEnabled(False)
        self.open_export_button.clicked.connect(self.open_export_location)
        header_layout.addWidget(self.open_export_button)

        main_layout.addLayout(header_layout)

        self.result_status_label = QLabel("")
        self.result_status_label.setWordWrap(True)
        self.result_status_label.setStyleSheet(
            "color: #64748B; font-size: 14px; background: transparent; line-height: 1.5;"
        )
        main_layout.addWidget(self.result_status_label)

        trans_card = self._create_translation_card()
        main_layout.addWidget(trans_card, 3)

        chat_card = self._create_chat_card()
        main_layout.addWidget(chat_card, 1)

        layout.addWidget(main_content, 1)

        return page

    def _create_history_sidebar(self):
        card = QFrame()
        card.setObjectName("history_card")
        card.setFixedWidth(290)
        card.setStyleSheet(
            "QFrame#history_card { background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 24px; }"
        )

        layout = QVBoxLayout(card)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        header_row = QHBoxLayout()
        title = QLabel("历史会话")
        title.setStyleSheet("font-size: 16px; font-weight: 700; color: #0F172A; background: transparent;")
        header_row.addWidget(title)
        header_row.addStretch()

        self.history_refresh_button = QPushButton("刷新")
        self.history_refresh_button.setObjectName("rewrite_action_btn")
        self.history_refresh_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.history_refresh_button.setFixedHeight(32)
        self.history_refresh_button.clicked.connect(self.refresh_history_sidebar)
        header_row.addWidget(self.history_refresh_button)
        layout.addLayout(header_row)

        self.history_status_label = QLabel("已保存的会话会显示在这里。")
        self.history_status_label.setWordWrap(True)
        self.history_status_label.setStyleSheet(
            "color: #64748B; font-size: 12px; background: transparent; line-height: 1.4;"
        )
        layout.addWidget(self.history_status_label)

        self.history_search_input = QLineEdit()
        self.history_search_input.setPlaceholderText("搜索标题、链接、正文、问答...")
        self.history_search_input.setClearButtonEnabled(True)
        self.history_search_input.textChanged.connect(self.on_history_search_changed)
        self.history_search_input.setStyleSheet(
            "QLineEdit { background-color: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 14px; padding: 10px 12px; }"
        )
        layout.addWidget(self.history_search_input)

        self.history_list_widget = QListWidget()
        self.history_list_widget.setObjectName("history_list")
        self.history_list_widget.itemClicked.connect(self._on_history_item_clicked)
        self.history_list_widget.setSpacing(8)
        self.history_list_widget.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )
        self.history_list_widget.setStyleSheet(
            "QListWidget#history_list { background: transparent; border: none; }"
        )
        layout.addWidget(self.history_list_widget, 1)

        return card

    def go_to_input_page(self):
        self.stack.setCurrentIndex(0)
        self.reset_ui()
        self.clear_results()

    def _create_settings_card(self):
        card = QFrame()
        card.setObjectName("rewrite_card")
        card.setStyleSheet("QFrame#rewrite_card { background-color: #FFFFFF; border-radius: 16px; padding: 4px; }")
        card_layout = QVBoxLayout(card)
        card_layout.setSpacing(12)
        card_layout.setContentsMargins(16, 16, 16, 16)

        header = QLabel("翻译设置")
        header.setStyleSheet("font-size: 14px; font-weight: 600; color: #1D1D1F; background: transparent;")
        card_layout.addWidget(header)

        provider_row = QHBoxLayout()
        provider_row.setSpacing(8)
        provider_label = QLabel("翻译服务")
        provider_label.setStyleSheet("color: #86868B; font-size: 13px; background: transparent;")
        provider_label.setFixedWidth(70)
        self.provider_combo = QComboBox()
        self.provider_combo.addItems(["deepseek", "openai", "ollama", "lmstudio"])
        self.provider_combo.setCurrentText("deepseek")
        self.provider_combo.currentTextChanged.connect(lambda _text: self._persist_settings())
        provider_row.addWidget(provider_label)
        provider_row.addWidget(self.provider_combo, 1)
        card_layout.addLayout(provider_row)

        skill_row = QHBoxLayout()
        skill_row.setSpacing(8)
        skill_label = QLabel("写作风格")
        skill_label.setStyleSheet("color: #86868B; font-size: 13px; background: transparent;")
        skill_label.setFixedWidth(70)
        self.skill_combo = QComboBox()
        if hasattr(self, "skills_dict") and self.skills_dict:
            self._populate_skill_combo()
        else:
            self.skill_combo.addItem("默认内置风格（晚点通用）", BUILTIN_STYLE_KEY)
        self.skill_combo.setToolTip(
            "优先级：如果所选风格文件包含 {{transcript}}，将作为完整写作 Prompt 直接执行；"
            "否则回退到后端内置的改写参考与题材路由。"
        )
        self.skill_combo.currentIndexChanged.connect(lambda _index: self._persist_settings())
        skill_row.addWidget(skill_label)
        skill_row.addWidget(self.skill_combo, 1)
        card_layout.addLayout(skill_row)

        skill_hint = QLabel(
            "导入格式：把 .md 或 .txt 风格文件直接放进 skills/。"
            "包含 {{transcript}} 时按完整 prompt 执行；否则只作为写作风格提示。"
        )
        skill_hint.setWordWrap(True)
        skill_hint.setStyleSheet(
            "color: #64748B; font-size: 12px; line-height: 1.4; background: transparent;"
        )
        card_layout.addWidget(skill_hint)

        api_row = QHBoxLayout()
        api_row.setSpacing(8)
        api_label = QLabel("API Key")
        api_label.setStyleSheet("color: #86868B; font-size: 13px; background: transparent;")
        api_label.setFixedWidth(70)
        self.api_key_input = QLineEdit()
        self.api_key_input.setEchoMode(QLineEdit.EchoMode.Password)
        self.api_key_input.setPlaceholderText("粘贴 DeepSeek API Key")
        self.api_key_input.textChanged.connect(lambda _text: self._persist_settings())
        api_row.addWidget(api_label)
        api_row.addWidget(self.api_key_input, 1)
        card_layout.addLayout(api_row)

        model_row = QHBoxLayout()
        model_row.setSpacing(8)
        model_label = QLabel("模型")
        model_label.setStyleSheet("color: #86868B; font-size: 13px; background: transparent;")
        model_label.setFixedWidth(70)
        self.model_input = QLineEdit()
        self.model_input.setText("deepseek-chat")
        self.model_input.textChanged.connect(lambda _text: self._persist_settings())
        model_row.addWidget(model_label)
        model_row.addWidget(self.model_input, 1)
        card_layout.addLayout(model_row)

        log_button = QPushButton("查看运行日志")
        log_button.setObjectName("rewrite_action_btn")
        log_button.setCursor(Qt.CursorShape.PointingHandCursor)
        log_button.clicked.connect(self.show_runtime_logs)
        card_layout.addWidget(log_button)

        return card

    def _create_translation_card(self):
        card = QFrame()
        card.setObjectName("rewrite_card")
        
        card_layout = QVBoxLayout(card)
        card_layout.setSpacing(0)
        card_layout.setContentsMargins(0, 0, 0, 0)

        # Body
        body = QFrame()
        body.setObjectName("rewrite_body")
        body_layout = QVBoxLayout(body)
        body_layout.setContentsMargins(32, 32, 32, 32)
        body_layout.setSpacing(0)

        self.trans_text = QTextEdit()
        self.trans_text.setReadOnly(True)
        self.trans_text.setPlaceholderText("改写结果会显示在这里，支持复制和导出。")
        
        css = """
        h1 { font-size: 24px; font-weight: bold; color: #1D1D1F; margin-top: 16px; margin-bottom: 12px; }
        h2 { font-size: 20px; font-weight: bold; color: #1D1D1F; margin-top: 16px; margin-bottom: 12px; }
        h3 { font-size: 16px; font-weight: bold; color: #1D1D1F; margin-top: 12px; margin-bottom: 8px; }
        p { font-size: 15px; color: #334155; margin-top: 0px; margin-bottom: 12px; line-height: 1.6; }
        ul { margin-top: 0px; margin-bottom: 12px; }
        li { font-size: 15px; color: #334155; line-height: 1.6; }
        """
        self.trans_text.document().setDefaultStyleSheet(css)
        
        body_layout.addWidget(self.trans_text)

        self.rewrite_issue_label = QLabel("")
        self.rewrite_issue_label.setWordWrap(True)
        self.rewrite_issue_label.setVisible(False)
        self.rewrite_issue_label.setStyleSheet(
            "color: #B45309; font-size: 12px; background: transparent; margin-top: 10px;"
        )
        body_layout.addWidget(self.rewrite_issue_label)

        card_layout.addWidget(body, 1)

        return card

    def _create_chat_card(self):
        card = QFrame()
        card.setObjectName("chat_card")
        card_layout = QVBoxLayout(card)
        card_layout.setSpacing(12)
        card_layout.setContentsMargins(12, 12, 12, 12)

        self.chat_messages_scroll = QScrollArea()
        self.chat_messages_scroll.setWidgetResizable(True)
        self.chat_messages_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.chat_messages_scroll.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        )
        self.chat_messages_scroll.setStyleSheet(
            "QScrollArea { background: transparent; border: none; }"
        )

        self.chat_messages_widget = QWidget()
        self.chat_messages_layout = QVBoxLayout(self.chat_messages_widget)
        self.chat_messages_layout.setContentsMargins(0, 0, 0, 0)
        self.chat_messages_layout.setSpacing(12)

        self.chat_empty_state = QFrame()
        self.chat_empty_state.setObjectName("chat_empty_state")
        empty_layout = QVBoxLayout(self.chat_empty_state)
        empty_layout.setContentsMargins(20, 24, 20, 24)
        empty_layout.setSpacing(8)
        empty_label = QLabel("先运行一次转写和翻译，然后在这里继续追问。")
        empty_label.setWordWrap(True)
        empty_label.setStyleSheet(
            "color: #64748B; font-size: 13px; background: transparent;"
        )
        empty_layout.addWidget(empty_label)
        self.chat_messages_layout.addWidget(self.chat_empty_state)
        self.chat_messages_layout.addStretch(1)

        self.chat_messages_scroll.setWidget(self.chat_messages_widget)
        card_layout.addWidget(self.chat_messages_scroll, 1)

        composer_frame = QFrame()
        composer_frame.setObjectName("chat_composer")
        composer_layout = QHBoxLayout(composer_frame)
        composer_layout.setContentsMargins(20, 10, 10, 10)
        composer_layout.setSpacing(12)

        self.chat_input = ChatTextEdit()
        self.chat_input.setPlaceholderText("输入问题，回车发送，Shift+Enter 换行。")
        self.chat_input.setFixedHeight(80)
        self.chat_input.returnPressed.connect(self.send_chat)
        self.chat_input.setEnabled(False)
        composer_layout.addWidget(self.chat_input, 1)

        self.chat_button = QPushButton("↑")
        self.chat_button.setObjectName("chat_send_btn")
        self.chat_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.chat_button.clicked.connect(self.send_chat)
        self.chat_button.setEnabled(False)
        composer_layout.addWidget(self.chat_button)

        card_layout.addWidget(composer_frame)

        return card

    def _clear_chat_messages(self):
        if not hasattr(self, "chat_messages_layout"):
            return

        for widget in list(self.chat_message_widgets):
            try:
                self.chat_messages_layout.removeWidget(widget)
                widget.setParent(None)
                widget.deleteLater()
            except Exception:
                pass

        self.chat_message_widgets = []
        self.chat_pending_message_widget = None
        self.chat_pending_message_label = None
        if hasattr(self, "chat_empty_state"):
            self.chat_empty_state.setVisible(True)
        if hasattr(self, "chat_messages_widget"):
            self.chat_messages_widget.adjustSize()
        self._scroll_chat_to_bottom()

    def _scroll_chat_to_bottom(self):
        if not hasattr(self, "chat_messages_scroll"):
            return
        scroll_bar = self.chat_messages_scroll.verticalScrollBar()
        scroll_bar.setValue(scroll_bar.maximum())

    def _chat_bubble_style(self, role: str, pending: bool = False, is_error: bool = False) -> str:
        if is_error:
            return (
                "QFrame { background-color: #FEF2F2; border: 1px solid #FCA5A5; "
                "border-radius: 16px; }"
            )
        if role == "user":
            return (
                "QFrame { background-color: #DBEAFE; border: 1px solid #93C5FD; "
                "border-radius: 16px; }"
            )
        if pending:
            return (
                "QFrame { background-color: #F8FAFC; border: 1px solid #E2E8F0; "
                "border-radius: 16px; }"
            )
        return (
            "QFrame { background-color: #FFFFFF; border: 1px solid #E2E8F0; "
            "border-radius: 16px; }"
        )

    def _create_chat_message_widget(
        self,
        role: str,
        text: str,
        pending: bool = False,
        is_error: bool = False,
    ):
        widget = QFrame()
        widget.setStyleSheet(self._chat_bubble_style(role, pending=pending, is_error=is_error))
        layout = QVBoxLayout(widget)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(6)

        header = QLabel("我" if role == "user" else "内容助手")
        header.setStyleSheet("color: #64748B; font-size: 12px; background: transparent;")

        content = QLabel(text)
        content.setWordWrap(True)
        content.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        if is_error:
            content.setStyleSheet(
                "color: #B91C1C; font-size: 14px; line-height: 1.5; background: transparent;"
            )
        elif pending:
            content.setStyleSheet(
                "color: #64748B; font-size: 14px; line-height: 1.5; font-style: italic; background: transparent;"
            )
        else:
            content.setStyleSheet(
                "color: #0F172A; font-size: 14px; line-height: 1.5; background: transparent;"
            )

        layout.addWidget(header)
        layout.addWidget(content)
        widget._chat_content_label = content
        widget._chat_role = role
        widget._chat_pending = pending
        widget._chat_error = is_error
        return widget

    def _append_chat_message(self, role: str, text: str, pending: bool = False, is_error: bool = False):
        if hasattr(self, "chat_empty_state"):
            self.chat_empty_state.setVisible(False)

        widget = self._create_chat_message_widget(role, text, pending=pending, is_error=is_error)
        insert_index = max(0, self.chat_messages_layout.count() - 1)
        self.chat_messages_layout.insertWidget(insert_index, widget)
        self.chat_message_widgets.append(widget)
        self._scroll_chat_to_bottom()
        return widget

    def _set_pending_chat_message(self, text: str, is_error: bool = False):
        widget = self.chat_pending_message_widget
        if widget is None:
            self._append_chat_message(
                "assistant",
                text,
                pending=not is_error,
                is_error=is_error,
            )
            self.chat_pending_message_widget = None
            self.chat_pending_message_label = None
            return

        widget.setStyleSheet(self._chat_bubble_style("assistant", pending=False, is_error=is_error))
        label = getattr(widget, "_chat_content_label", None)
        if label is not None:
            label.setText(text)
            if is_error:
                label.setStyleSheet(
                    "color: #B91C1C; font-size: 14px; line-height: 1.5; background: transparent;"
                )
            else:
                label.setStyleSheet(
                    "color: #0F172A; font-size: 14px; line-height: 1.5; background: transparent;"
                )
        widget._chat_pending = False
        widget._chat_error = is_error
        self.chat_pending_message_widget = None
        self.chat_pending_message_label = None
        self._scroll_chat_to_bottom()

    def _ensure_backend_available(self) -> bool:
        if self.backend_url:
            try:
                resp = requests.get(f"{self.backend_url}{BACKEND_HEALTH_PATH}", timeout=2)
                if resp.status_code == 200:
                    return True
            except Exception:
                pass

        try:
            self.backend_url = resolve_backend_base_url(force_probe=True)
            return True
        except Exception as exc:
            self._backend_bootstrap_error = str(exc)
            return False

    def _is_valid_youtube_url(self, url: str) -> bool:
        normalized = str(url or "").strip()
        if not normalized:
            return False
        return bool(
            re.search(
                r"(youtube\.com/watch\?v=|youtu\.be/|youtube\.com/shorts/)",
                normalized,
                flags=re.IGNORECASE,
            )
        )

    def copy_rewrite_text(self):
        text = self._current_rewrite_text()
        if not text:
            QMessageBox.information(self, "复制", "当前没有可复制的改写内容。")
            return
        QApplication.clipboard().setText(text)
        self._set_status_text("已复制到剪贴板")

    def _extract_export_basename(self, text: str) -> str:
        title = self._current_video_title()
        if title:
            return self._sanitize_filename_fragment(title)

        heading = self._first_markdown_heading(text)
        if heading:
            return self._sanitize_filename_fragment(heading)

        first_line = self._first_non_empty_line(text)
        if first_line:
            return self._sanitize_filename_fragment(first_line)

        return ""

    def _first_markdown_heading(self, text: str) -> str:
        for line in str(text or "").splitlines():
            stripped = line.strip()
            if stripped.startswith("#"):
                candidate = stripped.lstrip("#").strip()
                if candidate:
                    return candidate
        return ""

    def _first_non_empty_line(self, text: str) -> str:
        for line in str(text or "").splitlines():
            stripped = line.strip()
            if stripped:
                return stripped
        return ""

    def _sanitize_filename_fragment(self, text: str) -> str:
        collapsed = " ".join(str(text or "").split())
        if not collapsed:
            return ""

        invalid_chars = set('<>:"/\\|?*')
        sanitized = "".join("_" if ch in invalid_chars else ch for ch in collapsed)
        sanitized = sanitized.strip(" ._")
        if len(sanitized) > 80:
            sanitized = sanitized[:80].rstrip(" ._")
        return sanitized

    def export_rewrite_text(self):
        text = self._current_rewrite_text()
        if not text:
            QMessageBox.information(self, "导出", "当前没有可导出的改写内容。")
            return

        export_basename = self._extract_export_basename(text)
        if export_basename:
            default_name = f"rewrite-{export_basename}.docx"
        else:
            default_name = f"rewrite-{datetime.now().strftime('%Y%m%d-%H%M%S')}.docx"
        file_path, selected_filter = QFileDialog.getSaveFileName(
            self,
            "导出改写内容",
            os.path.join(os.path.expanduser("~"), default_name),
            "Word Documents (*.docx);;Markdown Files (*.md);;Text Files (*.txt);;HTML Files (*.html);;JSON Files (*.json)",
        )
        if not file_path:
            return

        _, ext = os.path.splitext(file_path)
        if not ext:
            if "Word Documents" in selected_filter:
                ext = ".docx"
            elif "Text Files" in selected_filter:
                ext = ".txt"
            elif "HTML Files" in selected_filter:
                ext = ".html"
            elif "JSON Files" in selected_filter:
                ext = ".json"
            else:
                ext = ".md"
            file_path = f"{file_path}{ext}"

        try:
            if ext.lower() == ".docx":
                self._write_docx_export(file_path, text)
            else:
                payload = self._build_export_content(text, ext.lower())
                with open(file_path, "w", encoding="utf-8") as handle:
                    handle.write(payload)
        except Exception as exc:
            QMessageBox.critical(self, "导出失败", str(exc))
            return

        self.last_exported_file_path = file_path
        self.open_export_button.setEnabled(True)
        self._set_status_text(f"已导出: {os.path.basename(file_path)}")

    def open_export_location(self):
        file_path = str(self.last_exported_file_path or "").strip()
        if not file_path or not os.path.exists(file_path):
            QMessageBox.information(self, "打开位置", "还没有可打开的导出文件。")
            return
        try:
            system = platform.system()
            if system == "Darwin":
                subprocess.run(["open", "-R", file_path], check=False)
            elif system == "Windows":
                subprocess.run(["explorer", "/select,", file_path], check=False)
            else:
                subprocess.run(["xdg-open", os.path.dirname(file_path)], check=False)
        except Exception as exc:
            QMessageBox.warning(self, "打开位置失败", str(exc))

    def show_runtime_logs(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("运行日志")
        dialog.resize(920, 680)

        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        log_view = QPlainTextEdit()
        log_view.setReadOnly(True)
        log_view.setPlaceholderText("这里会显示 backend / desktop 的最近运行日志。")
        log_view.setPlainText(self._collect_runtime_logs())
        log_view.setStyleSheet(
            "QPlainTextEdit { background-color: #0F172A; color: #E2E8F0; border-radius: 16px; padding: 12px; font-family: Menlo, Monaco, monospace; font-size: 12px; }"
        )
        layout.addWidget(log_view, 1)

        close_button = QPushButton("关闭")
        close_button.clicked.connect(dialog.accept)
        layout.addWidget(close_button, 0, Qt.AlignmentFlag.AlignRight)

        dialog.exec()

    def _collect_runtime_logs(self) -> str:
        root_dir = os.path.dirname(os.path.abspath(__file__))
        candidates = [
            os.path.join(root_dir, "backend_run.log"),
            os.path.join(root_dir, "desktop_run.log"),
            os.path.join(root_dir, "desktop_crash.log"),
        ]
        parts = []
        for path in candidates:
            label = os.path.basename(path)
            parts.append(f"===== {label} =====")
            parts.append(self._read_log_tail(path))
            parts.append("")
        return "\n".join(parts).strip()

    def _read_log_tail(self, path: str, max_lines: int = 120) -> str:
        if not os.path.exists(path):
            return "(文件不存在)"
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                lines = handle.readlines()
        except Exception as exc:
            return f"(读取失败: {exc})"
        if len(lines) > max_lines:
            lines = lines[-max_lines:]
        return "".join(lines).strip() or "(空)"

    def _build_export_content(self, text: str, extension: str) -> str:
        if extension == ".txt":
            return text
        if extension == ".html":
            paragraphs = "</p><p>".join(
                html.escape(item) for item in text.split("\n\n") if item.strip()
            )
            return (
                "<!doctype html><html><head><meta charset=\"utf-8\">"
                "<title>中文改写</title></head><body><h1>中文改写</h1>"
                f"<p>{paragraphs}</p></body></html>"
            )
        if extension == ".json":
            return json.dumps(
                {
                    "title": "中文改写",
                    "source_url": self.url_input.text().strip(),
                    "exported_at": datetime.now().isoformat(timespec="seconds"),
                    "rewritten_text": text,
                },
                ensure_ascii=False,
                indent=2,
            )
        return f"# 中文改写\n\n{text}\n"

    def _write_docx_export(self, file_path: str, text: str) -> None:
        try:
            from docx import Document
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
            from docx.shared import Pt
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx 依赖，请重新运行桌面启动脚本安装依赖。") from exc

        document = Document()
        document.core_properties.title = "中文改写"
        document.core_properties.subject = "Translation Writing Workbench export"
        document.core_properties.author = "Translation Writing Workbench"

        title = self._current_video_title()
        heading_text = title or "中文改写"
        heading = document.add_heading(heading_text, level=0)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        meta_bits = []
        video_url = ""
        if self.current_content:
            video = self.current_content.get("video")
            if isinstance(video, dict):
                video_url = str(video.get("url", "") or "").strip()
            content_context_id = str(
                self.current_content.get("content_context_id", "")
            ).strip()
            if video_url:
                meta_bits.append(f"来源: {video_url}")
            if content_context_id:
                meta_bits.append(f"会话: {content_context_id}")
        meta_bits.append(f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        for bit in meta_bits:
            paragraph = document.add_paragraph(bit)
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        normal_style = document.styles["Normal"]
        normal_style.font.name = "PingFang SC"
        normal_style.font.size = Pt(11)

        self._append_docx_body(document, text)
        document.save(file_path)

    def _append_docx_body(self, document, text: str) -> None:
        buffer: list[str] = []

        def flush_buffer():
            if not buffer:
                return
            paragraph_text = " ".join(
                part.strip() for part in buffer if part.strip()
            ).strip()
            buffer.clear()
            if paragraph_text:
                document.add_paragraph(paragraph_text)

        for raw_line in str(text or "").splitlines():
            line = raw_line.rstrip()
            stripped = line.strip()
            if not stripped:
                flush_buffer()
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if heading_match:
                flush_buffer()
                level = min(len(heading_match.group(1)), 4)
                document.add_heading(heading_match.group(2).strip(), level=level)
                continue

            bullet_match = re.match(r"^[-*•]\s+(.*)$", stripped)
            if bullet_match:
                flush_buffer()
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(bullet_match.group(1).strip())
                continue

            ordered_match = re.match(r"^\d+[.)]\s+(.*)$", stripped)
            if ordered_match:
                flush_buffer()
                paragraph = document.add_paragraph(style="List Number")
                paragraph.add_run(ordered_match.group(1).strip())
                continue

            buffer.append(stripped)

        flush_buffer()

    def check_backend(self):
        if self.backend_bootstrap_worker and self.backend_bootstrap_worker.isRunning():
            return

        self._set_status_text("正在检查后端...")
        self.backend_bootstrap_worker = BackendBootstrapWorker(self)
        self.backend_bootstrap_worker.result.connect(self.on_backend_bootstrap_result)
        self.backend_bootstrap_worker.finished.connect(self._on_backend_bootstrap_finished)
        self.backend_bootstrap_worker.finished.connect(self.backend_bootstrap_worker.deleteLater)
        self.backend_bootstrap_worker.start()

    def on_backend_bootstrap_result(self, success: bool, backend_url: str, info: str):
        self._backend_bootstrap_error = str(info or "")
        if success and backend_url:
            self.backend_url = backend_url
            self._set_backend_online_ui()
            self._backend_bootstrap_error = ""
            self.refresh_history_sidebar()
            self._set_status_text("后端已连接")
            return

        self.backend_url = None
        self._set_backend_offline_ui()
        self._set_status_text(self._backend_bootstrap_error or "后端未连接")

    def _on_backend_bootstrap_finished(self):
        self.backend_bootstrap_worker = None

    def start_processing(self):
        if not self._ensure_backend_available():
            extra = (
                f"\n\n自动启动信息：{self._backend_bootstrap_error}"
                if self._backend_bootstrap_error
                else ""
            )
            QMessageBox.critical(self, "错误", f"后端未连接，且自动拉起失败。{extra}")
            return

        url = self.url_input.text().strip()
        if not url:
            QMessageBox.critical(self, "错误", "请输入 YouTube 链接")
            return
        if not self._is_valid_youtube_url(url):
            QMessageBox.critical(self, "错误", "请输入有效的 YouTube 链接")
            return

        self.run_button.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self._set_status_text("处理中...")
        self.clear_results()
        self._clear_chat_messages()

        selected_skill_key = str(self.skill_combo.currentData() or BUILTIN_STYLE_KEY).strip()
        skill_prompt = self.skills_dict.get(selected_skill_key, "")
        if selected_skill_key and selected_skill_key != BUILTIN_STYLE_KEY:
            validation = self.skill_prompt_validation.get(selected_skill_key) or validate_rewrite_prompt(skill_prompt)
            if not validation.is_valid:
                QMessageBox.warning(self, "写作风格无效", "\n".join(validation.errors))
                self.run_button.setEnabled(True)
                self.progress_bar.setVisible(False)
                self._set_status_text("")
                return
        
        self.worker = WorkerThread(
            backend_url=self.backend_url,
            url=url,
            provider=self.provider_combo.currentText(),
            source_mode="subtitle_first",
            api_key=self.api_key_input.text().strip(),
            model=self.model_input.text().strip(),
            skill_prompt=skill_prompt,
        )
        self.worker.progress.connect(self.update_progress)
        self.worker.result.connect(self.display_results)
        self.worker.error.connect(self.show_error)
        self.worker.finished.connect(self.reset_ui)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.start()

    def update_progress(self, value, text):
        self.progress_bar.setValue(value)
        self._set_status_text(text)

    def display_results(self, data):
        rewritten_text = str(data.get("rewrite_zh", "")).strip()
        self.current_rewrite_raw_text = rewritten_text
        self.current_rewrite_source_text = str(data.get("rewrite_source_text", "")).strip()
        self.current_rewrite_focus = str(data.get("rewrite_focus", "")).strip()
        self._update_rewrite_quality_banner(data.get("rewrite_quality_issues", []))
        self.last_exported_file_path = ""
        self.open_export_button.setEnabled(False)
        self._clear_chat_messages()
        if rewritten_text:
            self.trans_text.setMarkdown(rewritten_text)
        else:
            translation = data.get("translation_zh", {})
            zh_segments = translation.get("segments", [])
            if zh_segments:
                paragraphs = []
                for seg in zh_segments:
                    text = str(seg.get("translated_text", "")).strip()
                    if text:
                        paragraphs.append(text)
                self.trans_text.setMarkdown("\n\n".join(paragraphs))
            else:
                self.trans_text.setMarkdown("暂无可展示内容")

        has_rewrite = bool(self._current_rewrite_text())
        self.copy_button.setEnabled(has_rewrite)
        self.export_button.setEnabled(has_rewrite)

        self.current_content = data
        self.chat_input.setEnabled(True)
        self.chat_button.setEnabled(True)
        self.rewrite_again_button.setEnabled(bool(rewritten_text or self.current_rewrite_source_text))
        self.current_history_id = str(data.get("content_context_id", "")).strip()
        self.stack.setCurrentIndex(1)
        self.refresh_history_sidebar()

    def send_chat(self):
        if not self._ensure_backend_available():
            QMessageBox.warning(self, "提示", "后端未连接，无法发送提问")
            return

        if not self.current_content:
            QMessageBox.warning(self, "提示", "请先处理一个视频")
            return

        question = self.chat_input.toPlainText().strip()
        if not question:
            return

        self.chat_input.clear()
        self._append_chat_message("user", question)
        pending_widget = self._append_chat_message("assistant", "思考中...", pending=True)
        self.chat_pending_message_widget = pending_widget
        self.chat_pending_message_label = getattr(pending_widget, "_chat_content_label", None)

        content_context_id = self.current_content.get("content_context_id", "")
        transcript_en = self.current_content.get("transcript_en", {}).get("text", "")
        translation_zh = "\n".join(
            seg.get("translated_text", "")
            for seg in self.current_content.get("translation_zh", {}).get("segments", [])
        )

        self.chat_button.setEnabled(False)
        self.chat_button.setText("…")

        chat_config = {
            "provider": self.provider_combo.currentText(),
        }
        api_key = self.api_key_input.text().strip()
        model = self.model_input.text().strip()
        if api_key:
            chat_config["api_key"] = api_key
        if model:
            chat_config["model"] = model

        self.chat_worker = ChatWorker(
            self.backend_url,
            question,
            content_context_id,
            transcript_en,
            translation_zh,
            chat_config,
        )
        self.chat_worker.response.connect(self.on_chat_response)
        self.chat_worker.error.connect(self.on_chat_error)
        self.chat_worker.finished.connect(self._on_chat_finished)
        self.chat_worker.finished.connect(self.chat_worker.deleteLater)
        self.chat_worker.start()

    def rewrite_current_result(self):
        if not self._ensure_backend_available():
            QMessageBox.warning(self, "提示", "后端未连接，无法重新改写")
            return

        source_text = self._current_rewrite_source_text()
        if not source_text:
            QMessageBox.warning(self, "提示", "当前没有可重新改写的原文。")
            return

        rewrite_focus = self._current_rewrite_focus()
        if not rewrite_focus:
            QMessageBox.warning(self, "提示", "当前没有可用的写作风格。")
            return

        translation_config = {
            "provider": self.provider_combo.currentText(),
        }
        api_key = self.api_key_input.text().strip()
        model = self.model_input.text().strip()
        if api_key:
            translation_config["api_key"] = api_key
        if model:
            translation_config["model"] = model

        self.rewrite_again_button.setEnabled(False)
        self.rewrite_again_button.setText("改写中…")
        self._set_status_text("正在重新改写...")
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(35)

        content_context_id = str(self.current_content.get("content_context_id", "")).strip() if self.current_content else ""
        self.rewrite_worker = RewriteWorker(
            self.backend_url,
            source_text,
            rewrite_focus,
            translation_config,
            content_context_id,
        )
        self.rewrite_worker.progress.connect(self.update_progress)
        self.rewrite_worker.result.connect(self.on_rewrite_again_result)
        self.rewrite_worker.error.connect(self.on_rewrite_again_error)
        self.rewrite_worker.finished.connect(self._on_rewrite_again_finished)
        self.rewrite_worker.finished.connect(self.rewrite_worker.deleteLater)
        self.rewrite_worker.start()

    def on_rewrite_again_result(self, data):
        rewritten_text = str(data.get("rewrite_zh", "")).strip()
        if rewritten_text:
            self.current_rewrite_raw_text = rewritten_text
            self.current_rewrite_source_text = str(data.get("rewrite_source_text", "")).strip()
            self.current_rewrite_focus = str(data.get("rewrite_focus", "")).strip()
            self._update_rewrite_quality_banner(data.get("rewrite_quality_issues", []))
            self.trans_text.setMarkdown(rewritten_text)
            if self.current_content:
                self.current_content["rewrite_zh"] = rewritten_text
                self.current_content["rewrite_source_text"] = self.current_rewrite_source_text
                self.current_content["rewrite_focus"] = self.current_rewrite_focus
                self.current_content["rewrite_quality_issues"] = list(self.current_rewrite_quality_issues)
                self.current_content["rewrite_provider"] = str(data.get("rewrite_provider", "")).strip()
                self.current_content["rewrite_model"] = str(data.get("rewrite_model", "")).strip()
            self.copy_button.setEnabled(True)
            self.export_button.setEnabled(True)
            self.rewrite_again_button.setEnabled(True)
            self._set_status_text("重新改写完成")
            self.refresh_history_sidebar()
        else:
            QMessageBox.warning(self, "重新改写", "没有返回可展示的改写结果。")

    def on_rewrite_again_error(self, error):
        friendly = self._friendly_error_text(error)
        QMessageBox.warning(self, "重新改写失败", friendly)

    def _on_rewrite_again_finished(self):
        self.rewrite_again_button.setEnabled(True)
        self.rewrite_again_button.setText("再改写")
        self.progress_bar.setVisible(False)
        self.rewrite_worker = None

    def on_chat_response(self, answer):
        message = str(answer or "").strip() or "没有返回可展示的回答。"
        if self.chat_pending_message_widget is not None:
            self._set_pending_chat_message(message)
        else:
            self._append_chat_message("assistant", message)

    def on_chat_error(self, error):
        friendly = self._friendly_error_text(error)
        lowered = friendly.lower()
        if "authentication fails" in lowered or "invalid api key" in lowered or "incorrect api key" in lowered:
            friendly = "鉴权失败：当前 API Key 无效，请检查左侧 API Key 与 provider 是否匹配。"
        else:
            friendly = f"对话失败：{friendly}"
        if self.chat_pending_message_widget is not None:
            self._set_pending_chat_message(friendly, is_error=True)
        else:
            self._append_chat_message("assistant", friendly, is_error=True)

    def _on_chat_finished(self):
        self.chat_button.setEnabled(True)
        self.chat_button.setText("↑")
        self.chat_worker = None
        self.refresh_history_sidebar()

    def show_error(self, message):
        self._set_backend_offline_ui()
        friendly = self._friendly_error_text(message)
        self._set_status_text(f"错误: {friendly[:50]}")
        QMessageBox.critical(self, "处理错误", friendly)

    def reset_ui(self):
        self.run_button.setEnabled(True)
        self.progress_bar.setVisible(False)
        self._set_status_text("")
        self.worker = None

    def clear_results(self):
        self.trans_text.clear()
        self.current_content = None
        self.current_history_id = ""
        self.current_rewrite_raw_text = ""
        self.current_rewrite_source_text = ""
        self.current_rewrite_focus = ""
        self.current_rewrite_quality_issues = []
        self.last_exported_file_path = ""
        self.rewrite_issue_label.setVisible(False)
        self.rewrite_issue_label.setText("")
        self.chat_input.setEnabled(False)
        self.chat_button.setEnabled(False)
        self.rewrite_again_button.setEnabled(False)
        self.chat_input.clear()
        self.copy_button.setEnabled(False)
        self.export_button.setEnabled(False)
        self.open_export_button.setEnabled(False)
        self._clear_chat_messages()

    def closeEvent(self, event: QCloseEvent):
        still_running = False
        for thread_name in (
            "worker",
            "chat_worker",
            "rewrite_worker",
            "history_worker",
            "history_detail_worker",
        ):
            thread = getattr(self, thread_name, None)
            if thread is None:
                continue
            try:
                if thread.isRunning():
                    thread.requestInterruption()
                    thread.quit()
                    if not thread.wait(2000):
                        still_running = True
            except Exception:
                pass
            if not still_running:
                setattr(self, thread_name, None)

        if still_running:
            QMessageBox.warning(self, "请稍候", "后台任务仍在结束中，请稍后再关闭。")
            event.ignore()
            return

        super().closeEvent(event)


def main():
    _install_fault_handler()
    _install_global_exception_hook()
    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    window = YouTubeTranslatorApp()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
